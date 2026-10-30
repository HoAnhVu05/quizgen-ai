import io
from typing import List, Dict, Any
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

class WordExamExporter:
    @staticmethod
    def generate_docx(
        exam_title: str,
        subject: str,
        exam_code: str,
        questions: List[Dict[str, Any]]
    ) -> bytes:
        doc = Document()

        # Page margins
        for section in doc.sections:
            section.top_margin = Inches(0.75)
            section.bottom_margin = Inches(0.75)
            section.left_margin = Inches(0.75)
            section.right_margin = Inches(0.75)

        # Header 2 columns table
        header_table = doc.add_table(rows=1, cols=2)
        header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell_left, cell_right = header_table.rows[0].cells
        
        # Left header: University info
        p_left = cell_left.paragraphs[0]
        r1 = p_left.add_run("ĐẠI HỌC QUỐC GIA\n")
        r1.bold = True
        r1.font.size = Pt(11)
        r2 = p_left.add_run("KHOA KINH TẾ & QUẢN TRỊ\n-------------------")
        r2.font.size = Pt(10)

        # Right header: Exam details
        p_right = cell_right.paragraphs[0]
        p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r3 = p_right.add_run(f"ĐỀ THI KẾT THÚC HỌC PHẦN\nMôn thi: {subject}\n")
        r3.bold = True
        r3.font.size = Pt(11)
        r4 = p_right.add_run(f"Mã đề thi: {exam_code} | Thời gian: 60 phút")
        r4.font.size = Pt(10)

        doc.add_paragraph()  # Spacing

        # Title
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_title = p_title.add_run(exam_title.upper())
        r_title.bold = True
        r_title.font.size = Pt(14)

        doc.add_paragraph("-------------------------------------------------------------------------------------------------------")

        # Questions list
        answer_keys = []
        for idx, q in enumerate(questions, 1):
            p_q = doc.add_paragraph()
            p_q.paragraph_format.space_before = Pt(8)
            p_q.paragraph_format.space_after = Pt(4)
            
            r_num = p_q.add_run(f"Câu {idx}: ")
            r_num.bold = True
            p_q.add_run(q.get("question", ""))

            # Options
            options = q.get("options", [])
            for opt in options:
                p_opt = doc.add_paragraph()
                p_opt.paragraph_format.left_indent = Inches(0.3)
                p_opt.paragraph_format.space_after = Pt(2)
                p_opt.add_run(f"{opt.get('label')}. {opt.get('text')}")

            # Store correct answer
            correct_id = q.get("correct_answer_id") or q.get("correctAnswerId")
            correct_opt = next((o for o in options if o.get("id") == correct_id), None)
            answer_keys.append((idx, correct_opt.get("label") if correct_opt else "A"))

        # Add Answer Key Matrix at the end
        doc.add_page_break()
        p_ans_title = doc.add_paragraph()
        p_ans_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_ans = p_ans_title.add_run(f"BẢNG ĐÁP ÁN CHUẨN - MÃ ĐỀ {exam_code}")
        r_ans.bold = True
        r_ans.font.size = Pt(13)

        # Answer grid table (5 columns)
        ans_table = doc.add_table(rows=1, cols=5)
        ans_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_cells = ans_table.rows[0].cells
        for i in range(5):
            hdr_cells[i].text = "Câu - ĐA"

        row_cells = None
        for i, (q_idx, ans_label) in enumerate(answer_keys):
            col_idx = i % 5
            if col_idx == 0:
                row_cells = ans_table.add_row().cells
            row_cells[col_idx].text = f"Câu {q_idx}: {ans_label}"

        # Save to buffer
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
